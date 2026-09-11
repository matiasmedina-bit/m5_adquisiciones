"""
Carga del catálogo de un proveedor desde una planilla Excel (.xlsx).

Los proveedores de M5 envían sus listas de precios en Excel, con encabezados
que nunca son iguales entre uno y otro. Este módulo acepta esa variedad: busca
la fila de encabezados en las primeras filas de la hoja y reconoce cada columna
por sinónimos, sin importar mayúsculas, tildes ni espacios sobrantes.

Columnas reconocidas (basta con 'código' y 'descripción'):
    código      -> codigo, cod, sku, referencia, ref, item
    descripción -> descripcion, detalle, material, producto, nombre, glosa
    unidad      -> unidad de medida, um, u.m., medida
    precio      -> valor, valor unitario, precio unitario, neto, $
    condición   -> condicion de pago, forma de pago, pago
"""
from decimal import Decimal, InvalidOperation

from .models import Proveedor, ProveedorMaterial


# --------------------------------------------------------------------------
# Reconocimiento de encabezados
# --------------------------------------------------------------------------
SINONIMOS = {
    "codigo": ("codigo", "cod", "sku", "referencia", "ref", "item", "codigo proveedor",
               "codigo del proveedor", "n", "nro"),
    "descripcion": ("descripcion", "detalle", "material", "producto", "nombre",
                    "glosa", "articulo", "descripcion del material"),
    "unidad": ("unidad", "unidad de medida", "um", "u m", "medida", "un"),
    "precio": ("precio", "valor", "valor unitario", "precio unitario", "neto",
               "precio neto", "valor neto", "$", "monto"),
    "condicion": ("condicion", "condicion de pago", "forma de pago", "pago",
                  "condiciones", "condiciones de pago"),
}

_TILDES = str.maketrans("áéíóúüñÁÉÍÓÚÜÑ", "aeiouunAEIOUUN")

CONDICIONES = {
    "contado": "CONTADO",
    "al contado": "CONTADO",
    "efectivo": "CONTADO",
    "0": "CONTADO",
    "30": "30_DIAS",
    "30 dias": "30_DIAS",
    "30 d": "30_DIAS",
    "60": "60_DIAS",
    "60 dias": "60_DIAS",
    "60 d": "60_DIAS",
    "90": "90_DIAS",
    "90 dias": "90_DIAS",
    "90 d": "90_DIAS",
}


def _normalizar(texto):
    """minúsculas, sin tildes, sin puntuación de separación, sin espacios dobles."""
    if texto is None:
        return ""
    txt = str(texto).translate(_TILDES).lower().strip()
    for basura in (".", ":", "_", "-", "/", "(", ")"):
        txt = txt.replace(basura, " ")
    return " ".join(txt.split())


def _detectar_columnas(fila):
    """Devuelve {campo: índice} para los encabezados que reconoce en la fila."""
    encontrados = {}
    for idx, celda in enumerate(fila):
        clave = _normalizar(celda)
        if not clave:
            continue
        for campo, opciones in SINONIMOS.items():
            if campo in encontrados:
                continue
            if clave in opciones:
                encontrados[campo] = idx
                break
    return encontrados


def _a_entero(valor):
    """
    Convierte a entero CLP. Tolera '$ 1.234', '1.234', '1234,00', 1234.0.
    En Chile el punto es separador de miles, así que se descarta.
    """
    if valor is None or valor == "":
        return 0
    if isinstance(valor, (int, float, Decimal)):
        return int(round(float(valor)))
    txt = str(valor).strip().replace("$", "").replace(" ", "")
    if not txt:
        return 0
    if "," in txt:               # la coma es el decimal: 1.234,50
        txt = txt.replace(".", "").replace(",", ".")
    else:                        # sólo puntos: son miles
        txt = txt.replace(".", "")
    try:
        return int(round(float(Decimal(txt))))
    except (InvalidOperation, ValueError):
        return 0


def _a_condicion(valor):
    """Mapea el texto de la planilla a una de las opciones del modelo."""
    clave = _normalizar(valor)
    if not clave:
        return ""
    if clave in CONDICIONES:
        return CONDICIONES[clave]
    # '30 dias corridos', 'a 60 dias', 'pago 90 dias'...
    for numero, opcion in (("30", "30_DIAS"), ("60", "60_DIAS"), ("90", "90_DIAS")):
        if numero in clave:
            return opcion
    if "contado" in clave or "efectivo" in clave:
        return "CONTADO"
    return ""


# --------------------------------------------------------------------------
# Carga
# --------------------------------------------------------------------------
class ResultadoCarga:
    """Resumen de lo que hizo la importación, para mostrárselo al usuario."""

    def __init__(self):
        self.creados = 0
        self.actualizados = 0
        self.omitidos = []      # [(nro_fila, motivo)]
        self.retirados = 0      # marcados no disponibles al reemplazar
        self.columnas = {}

    @property
    def total_procesados(self):
        return self.creados + self.actualizados

    def resumen(self):
        partes = []
        if self.creados:
            partes.append(f"{self.creados} material(es) nuevo(s)")
        if self.actualizados:
            partes.append(f"{self.actualizados} actualizado(s)")
        if self.retirados:
            partes.append(f"{self.retirados} marcado(s) como no disponible(s)")
        if not partes:
            partes.append("ningún material cargado")
        texto = "Catálogo procesado: " + ", ".join(partes) + "."
        if self.omitidos:
            texto += f" Se omitieron {len(self.omitidos)} fila(s)."
        return texto


class ErrorCatalogo(Exception):
    """La planilla no se puede leer o no tiene las columnas mínimas."""


def importar_catalogo(proveedor: Proveedor, archivo, reemplazar=False) -> ResultadoCarga:
    """
    Lee el .xlsx y crea o actualiza los ProveedorMaterial del proveedor.

    Los materiales se identifican por 'código' dentro del proveedor, igual que
    la restricción unique_together del modelo. Un código que ya existe se
    actualiza; nunca se borra nada, para no romper el historial de órdenes de
    compra (RF-06). Con reemplazar=True, los que no vienen en la planilla se
    marcan como no disponibles en vez de eliminarse.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover
        raise ErrorCatalogo(
            "Falta la librería 'openpyxl'. Ejecuta: pip install -r requirements.txt"
        )

    try:
        libro = load_workbook(archivo, data_only=True, read_only=True)
    except Exception as exc:
        raise ErrorCatalogo(
            f"No se pudo abrir la planilla ({exc}). Verifica que sea un archivo .xlsx válido."
        )

    hoja = libro.active
    resultado = ResultadoCarga()

    # 1. Buscar la fila de encabezados en las primeras 10 filas: muchas planillas
    #    de proveedor traen el logo y datos de contacto antes de la tabla.
    columnas = {}
    fila_encabezado = 0
    filas = []
    for nro, fila in enumerate(hoja.iter_rows(values_only=True), start=1):
        filas.append((nro, fila))
        if not columnas and nro <= 10:
            posibles = _detectar_columnas(fila)
            if "codigo" in posibles and "descripcion" in posibles:
                columnas = posibles
                fila_encabezado = nro

    if not columnas:
        raise ErrorCatalogo(
            "No se encontraron las columnas mínimas en la planilla. "
            "Debe tener al menos una columna «Código» y una «Descripción» "
            "en alguna de las primeras 10 filas. "
            "Puedes descargar la plantilla de ejemplo desde la ficha del proveedor."
        )

    resultado.columnas = columnas
    i_cod = columnas["codigo"]
    i_desc = columnas["descripcion"]
    i_uni = columnas.get("unidad")
    i_pre = columnas.get("precio")
    i_con = columnas.get("condicion")

    codigos_vistos = set()

    # 2. Procesar las filas de datos
    for nro, fila in filas:
        if nro <= fila_encabezado:
            continue
        if fila is None or all(c is None or str(c).strip() == "" for c in fila):
            continue

        def celda(indice):
            if indice is None or indice >= len(fila):
                return None
            return fila[indice]

        codigo = str(celda(i_cod) or "").strip()
        descripcion = str(celda(i_desc) or "").strip()

        if not codigo:
            resultado.omitidos.append((nro, "sin código"))
            continue
        if not descripcion:
            resultado.omitidos.append((nro, "sin descripción"))
            continue
        if len(codigo) > 40:
            resultado.omitidos.append((nro, "el código supera los 40 caracteres"))
            continue
        clave = codigo.lower()
        if clave in codigos_vistos:
            resultado.omitidos.append((nro, f"código «{codigo}» repetido en la planilla"))
            continue
        codigos_vistos.add(clave)

        datos = {
            "descripcion": descripcion[:200],
            "unidad_medida": (str(celda(i_uni) or "").strip() or "un")[:20],
            "precio": _a_entero(celda(i_pre)),
            "condicion_pago": _a_condicion(celda(i_con)),
            "disponible": True,
        }

        existente = ProveedorMaterial.objects.filter(
            proveedor=proveedor, codigo__iexact=codigo
        ).first()
        if existente:
            for campo, valor in datos.items():
                setattr(existente, campo, valor)
            existente.save()
            resultado.actualizados += 1
        else:
            ProveedorMaterial.objects.create(proveedor=proveedor, codigo=codigo, **datos)
            resultado.creados += 1

    libro.close()

    # 3. Reemplazo: lo que no vino en la planilla deja de ofrecerse (RF-06)
    if reemplazar and codigos_vistos:
        pendientes = proveedor.materiales.filter(disponible=True)
        for material in pendientes:
            if material.codigo.lower() not in codigos_vistos:
                material.disponible = False
                material.save(update_fields=["disponible", "modificado"])
                resultado.retirados += 1

    return resultado


# --------------------------------------------------------------------------
# Plantilla de ejemplo para entregarle al proveedor
# --------------------------------------------------------------------------
def generar_plantilla() -> bytes:
    """Devuelve un .xlsx vacío con los encabezados y dos filas de ejemplo."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Catálogo"

    encabezados = ["Código", "Descripción", "Unidad", "Precio", "Condición de pago"]
    hoja.append(encabezados)

    relleno = PatternFill("solid", fgColor="1F3B4D")
    for col, _ in enumerate(encabezados, start=1):
        celda = hoja.cell(row=1, column=col)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center")

    hoja.append(["CEM-001", "Cemento Portland 25 kg", "saco", 5490, "30 días"])
    hoja.append(["FIE-014", "Fierro estriado 8 mm x 6 m", "un", 3990, "Contado"])

    for col, ancho in zip("ABCDE", (14, 42, 10, 12, 20)):
        hoja.column_dimensions[col].width = ancho
    hoja.freeze_panes = "A2"

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()

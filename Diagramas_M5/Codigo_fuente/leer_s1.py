import openpyxl, json
wb = openpyxl.load_workbook('/home/claude/esf/Grupo_23_Tabla_de_Esfuerzo_Acumulada.xlsx', data_only=True)

ws = wb['Planificación']
tareas, grupo = [], None
for r in range(6, ws.max_row + 1):
    a = ws.cell(r, 1).value
    if a is None:
        continue
    a = str(a).strip()
    if a.startswith('SPRINT 2'):
        break
    if a.startswith('TA '):
        grupo = a
        continue
    if not a.startswith('S1-'):
        continue
    tareas.append(dict(id=a, grupo=grupo,
                       desc=str(ws.cell(r, 2).value).strip(),
                       resp=str(ws.cell(r, 3).value).strip(),
                       est=ws.cell(r, 5).value,
                       ini=str(ws.cell(r, 6).value).strip(),
                       fin=str(ws.cell(r, 7).value).strip(),
                       real=ws.cell(r, 8).value,
                       estado=str(ws.cell(r, 9).value).strip()))
print('tareas sprint 1:', len(tareas))
print('HH est', sum(t['est'] for t in tareas), '| HH real', sum(t['real'] for t in tareas))
print('grupos:', sorted({t['grupo'] for t in tareas}))
for t in tareas[:3]:
    print('  ', t)

# serie de esfuerzo del sprint 1 (21 dias) desde Graficas acumuladas
wg = wb['Gráficas acumuladas']
serie = []
for r in range(5, wg.max_row + 1):
    if wg.cell(r, 3).value != 'Sprint 1':
        continue
    serie.append(dict(dia=wg.cell(r, 1).value, fecha=str(wg.cell(r, 2).value),
                      pend_real=wg.cell(r, 4).value, pend_plan=wg.cell(r, 5).value,
                      acum_real=wg.cell(r, 6).value, acum_plan=wg.cell(r, 7).value))
print('dias sprint 1 en graficas acumuladas:', len(serie))
print('  primero', serie[0]); print('  ultimo', serie[-1])

json.dump(dict(tareas=tareas, serie=serie), open('/home/claude/final/sprint1.json', 'w'),
          ensure_ascii=False, indent=1)

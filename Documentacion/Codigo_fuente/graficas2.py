# -*- coding: utf-8 -*-
"""Figuras del Incremento 2 generadas desde el Excel de esfuerzo acumulado."""
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
import openpyxl, os

XLS = '/home/claude/esf/Grupo_23_Tabla_de_Esfuerzo_Acumulada.xlsx'
OUT = '/home/claude/final/img'
os.makedirs(OUT, exist_ok=True)
plt.rcParams['font.family'] = 'DejaVu Sans'

MORADO, AZUL, GRIS = '#7030A0', '#1F3864', '#808080'
MORADO_CLARO, AZUL_CLARO = '#EFE4F8', '#D9E2F3'

wb = openpyxl.load_workbook(XLS, data_only=True)
wp, we, wg2, wga = wb['Planificación'], wb['Tabla de Esfuerzo'], wb['Gráficas Incremento 2'], wb['Gráficas acumuladas']

# ---------------------------------------------------------------- lectura
plan = {'1': [], '2': []}
grupos = {'1': [], '2': []}
actual = None
for r in range(5, wp.max_row + 1):
    a, b, c = wp.cell(r, 1).value, wp.cell(r, 2).value, wp.cell(r, 3).value
    if a and str(a).startswith('SPRINT 1'):
        actual = '1'; continue
    if a and str(a).startswith('SPRINT 2'):
        actual = '2'; continue
    if a and str(a).startswith('TA ') and not c:
        grupos[actual].append((len(plan[actual]), a)); continue
    if a and c:
        plan[actual].append((a, b, c, wp.cell(r, 5).value, wp.cell(r, 6).value,
                             wp.cell(r, 7).value, wp.cell(r, 8).value))

TOT1_EST = sum(t[3] for t in plan['1']); TOT1_REAL = sum(t[6] for t in plan['1'])
TOT2_EST = sum(t[3] for t in plan['2']); TOT2_REAL = sum(t[6] for t in plan['2'])
TOTA_EST, TOTA_REAL = TOT1_EST + TOT2_EST, TOT1_REAL + TOT2_REAL
print('Sprint 1: %d tareas %d/%d · Sprint 2: %d tareas %d/%d · acumulado %d/%d'
      % (len(plan['1']), TOT1_EST, TOT1_REAL, len(plan['2']), TOT2_EST, TOT2_REAL, TOTA_EST, TOTA_REAL))

g2 = [[wg2.cell(r, c).value for c in range(1, 7)] for r in range(5, 27)]
ga = [[wga.cell(r, c).value for c in range(1, 9)] for r in range(5, 48)]


# ---------------------------------------------------------------- tablas
def tabla_png(ruta, columnas, anchos, filas, alto_fila=0.28, fs=8.0, ancho=9.6):
    n = len(filas) + 1
    fig = plt.figure(figsize=(ancho, n * alto_fila + 0.25), dpi=200)
    ax = fig.add_axes([0, 0, 1, 1]); ax.axis('off'); ax.set_xlim(0, 1); ax.set_ylim(0, 1)
    total = sum(anchos); xs, acc = [], 0.0
    for a in anchos:
        xs.append(acc / total); acc += a
    xs.append(1.0)
    h = 1.0 / n

    def celda(x0, x1, y0, fondo, texto, color='black', negrita=False, centrado=False):
        ax.add_patch(Rectangle((x0, y0), x1 - x0, h, facecolor=fondo, edgecolor='#BFBFBF',
                               linewidth=0.5, zorder=1))
        if texto != '':
            ax.text((x0 + x1) / 2 if centrado else x0 + 0.004, y0 + h / 2, texto,
                    ha='center' if centrado else 'left', va='center', fontsize=fs, color=color,
                    fontweight='bold' if negrita else 'normal', zorder=2)

    y = 1.0 - h
    for j, c in enumerate(columnas):
        celda(xs[j], xs[j + 1], y, '#404040', c, 'white', True, True)
    for f in filas:
        y -= h
        if f['tipo'] == 'seccion':
            celda(0, 1, y, f['color'], f['celdas'][0], 'white', True)
        else:
            neg = f['tipo'] == 'total'
            for j, c in enumerate(f['celdas']):
                celda(xs[j], xs[j + 1], y, f.get('fondo', 'white'), str(c), 'black', neg,
                      j >= len(columnas) - f.get('ncentro', 3))
    fig.savefig(ruta, bbox_inches='tight', pad_inches=0.02, facecolor='white')
    plt.close(fig); print('->', ruta)


def corta(t, n):
    return t if len(t) <= n else t[:n - 1] + '…'


def tabla_planificacion(sprint, ruta, titulo, color, claro):
    filas = [{'tipo': 'seccion', 'celdas': [titulo], 'color': color}]
    grp = dict(grupos[sprint])
    for i, t in enumerate(plan[sprint]):
        if i in grp:
            filas.append({'tipo': 'seccion', 'celdas': [grp[i]], 'color': color})
        filas.append({'tipo': 'dato', 'fondo': claro, 'ncentro': 5,
                      'celdas': [t[0], corta(t[1], 68), t[2], t[3], t[4], t[5], t[6], 'Finalizado']})
    est = sum(t[3] for t in plan[sprint]); real = sum(t[6] for t in plan[sprint])
    filas.append({'tipo': 'total', 'fondo': claro, 'ncentro': 5,
                  'celdas': ['', 'Subtotal · %d tareas' % len(plan[sprint]), '', est, '', '', real, '']})
    tabla_png(ruta, ['ID', 'Tarea', 'Responsable', 'HH est.', 'Inicio', 'Término', 'HH real', 'Estado'],
              [7, 44, 15, 7, 9, 9, 7, 9], filas, alto_fila=0.245, fs=7.2, ancho=9.7)


tabla_planificacion('1', OUT + '/Fig_2.2a_Planificacion_inc1.png',
                    'SPRINT 1 · INCREMENTO 1 — trabajo heredado (%d tareas · %d HH estimadas)'
                    % (len(plan['1']), TOT1_EST), MORADO, MORADO_CLARO)
tabla_planificacion('2', OUT + '/Fig_2.2b_Planificacion_inc2.png',
                    'SPRINT 2 · INCREMENTO 2 — trabajo nuevo (%d tareas · %d HH estimadas)'
                    % (len(plan['2']), TOT2_EST), AZUL, AZUL_CLARO)

# tabla de esfuerzo del Sprint 2
filas = [{'tipo': 'dato', 'fondo': AZUL_CLARO, 'ncentro': 5,
          'celdas': [('Día %d' % f[0]) if f[0] else 'Inicio', f[1], f[2], f[3], f[4], f[5]]} for f in g2]
tabla_png(OUT + '/Fig_6.1_Esfuerzo_sprint2.png',
          ['Día', 'Fecha', 'Pendiente real (HH)', 'Pendiente planificado (HH)',
           'Acumulado real (HH)', 'Acumulado planificado (HH)'],
          [10, 12, 19, 21, 19, 21], filas, alto_fila=0.27, fs=8.4, ancho=7.6)

# tabla de esfuerzo acumulada, en dos bloques lado a lado
mitad = 22
izq, der = ga[:mitad], ga[mitad:]
filas = []
for i in range(max(len(izq), len(der))):
    a = izq[i] if i < len(izq) else None
    b = der[i] if i < len(der) else None
    cel = ([('Día %d' % a[0]) if a[0] else 'Inicio', a[1], a[3], a[4], a[5], a[6]] if a else ['', '', '', '', '', ''])
    cel += ([('Día %d' % b[0]), b[1], b[3], b[4], b[5], b[6]] if b else ['', '', '', '', '', ''])
    filas.append({'tipo': 'dato', 'ncentro': 12,
                  'fondo': MORADO_CLARO if (a and a[2] == 'Sprint 1') else AZUL_CLARO, 'celdas': cel})
CAB = ['Día', 'Fecha', 'Pend. real', 'Pend. plan.', 'Acum. real', 'Acum. plan.']
tabla_png(OUT + '/Fig_6.2_Esfuerzo_acumulado.png', CAB + CAB,
          [8, 10, 12, 12, 12, 12] * 2, filas, alto_fila=0.26, fs=7.6, ancho=9.7)


# ---------------------------------------------------------------- gráficas
def base(titulo, xlabel, ylabel):
    fig, ax = plt.subplots(figsize=(9.0, 4.2), dpi=200)
    ax.set_title(titulo, fontsize=11, fontweight='bold', color='#333333', pad=10)
    ax.set_xlabel(xlabel, fontsize=8.5); ax.set_ylabel(ylabel, fontsize=8.5)
    ax.grid(True, linewidth=0.4, color='#DDDDDD'); ax.set_axisbelow(True)
    for s in ('top', 'right'):
        ax.spines[s].set_visible(False)
    ax.tick_params(labelsize=7)
    return fig, ax


etq2 = [f[1] for f in g2]
x2 = list(range(len(g2)))
fig, ax = base('Burn-Down del Incremento 2 — esfuerzo pendiente', 'Día de sprint (Sprint 2)',
               'HH pendientes del incremento')
ax.plot(x2, [f[3] for f in g2], color=GRIS, ls='--', lw=1.8, label='Pendiente planificado')
ax.plot(x2, [f[2] for f in g2], color=AZUL, lw=2.4, label='Pendiente real · Incremento 2')
ax.set_xticks(x2); ax.set_xticklabels(etq2, rotation=45, ha='right')
ax.legend(fontsize=8, frameon=False)
fig.savefig(OUT + '/Fig_6.3_Burndown_inc2.png', bbox_inches='tight', facecolor='white'); plt.close(fig)

fig, ax = base('Burn-Up del Incremento 2 — esfuerzo acumulado', 'Día de sprint (Sprint 2)',
               'HH acumuladas del incremento')
ax.plot(x2, [f[5] for f in g2], color=GRIS, ls='--', lw=1.8, label='Acumulado planificado')
ax.plot(x2, [f[4] for f in g2], color=AZUL, lw=2.4, label='Acumulado real · Incremento 2')
ax.axhline(TOT2_REAL, color='#C00000', lw=1.2, ls=':', label='Alcance comprometido (%d HH)' % TOT2_REAL)
ax.set_xticks(x2); ax.set_xticklabels(etq2, rotation=45, ha='right')
ax.legend(fontsize=8, frameon=False, loc='lower right')
fig.savefig(OUT + '/Fig_6.5_Burnup_inc2.png', bbox_inches='tight', facecolor='white'); plt.close(fig)

etqa = [f[1] for f in ga]
xa = list(range(len(ga)))
corte = 21   # índice del día 21 dentro de la lista (día 0 en la posición 0)
fig, ax = base('Burn-Down acumulado — Incrementos 1 y 2', 'Día de sprint (Sprint 1 + Sprint 2)',
               'HH pendientes del proyecto')
ax.plot(xa, [f[4] for f in ga], color=GRIS, ls='--', lw=1.8, label='Pendiente planificado')
ax.plot(xa[:corte + 1], [f[3] for f in ga[:corte + 1]], color=MORADO, lw=2.4,
        label='Pendiente real · Incremento 1 (heredado)')
ax.plot(xa[corte:], [f[3] for f in ga[corte:]], color=AZUL, lw=2.4, label='Pendiente real · Incremento 2')
ax.axvline(corte, color='#BBBBBB', lw=1, ls=':')
ax.text(corte + 0.4, 20, 'cierre del Incremento 1', fontsize=7.5, color='#666666')
ax.set_xticks(xa[::2]); ax.set_xticklabels(etqa[::2], rotation=45, ha='right')
ax.legend(fontsize=8, frameon=False)
fig.savefig(OUT + '/Fig_6.4_Burndown_acumulado.png', bbox_inches='tight', facecolor='white'); plt.close(fig)

fig, ax = base('Burn-Up acumulado — Incrementos 1 y 2', 'Día de sprint (Sprint 1 + Sprint 2)',
               'HH acumuladas del proyecto')
ax.plot(xa, [f[6] for f in ga], color=GRIS, ls='--', lw=1.8, label='Acumulado planificado')
ax.plot(xa[:corte + 1], [f[5] for f in ga[:corte + 1]], color=MORADO, lw=2.4,
        label='Acumulado real · Incremento 1 (heredado)')
ax.plot(xa[corte:], [f[5] for f in ga[corte:]], color=AZUL, lw=2.4, label='Acumulado real · Incremento 2')
ax.step(xa, [f[7] for f in ga], where='post', color='#C00000', lw=1.2, ls=':',
        label='Alcance comprometido')
ax.axvline(corte, color='#BBBBBB', lw=1, ls=':')
ax.set_xticks(xa[::2]); ax.set_xticklabels(etqa[::2], rotation=45, ha='right')
ax.legend(fontsize=8, frameon=False, loc='upper left')
fig.savefig(OUT + '/Fig_6.6_Burnup_acumulado.png', bbox_inches='tight', facecolor='white'); plt.close(fig)

print('DATOS PARA EL TEXTO: sprint2 %d tareas %d/%d · desviación %d HH (%.1f %%) · acumulado %d/%d'
      % (len(plan['2']), TOT2_EST, TOT2_REAL, TOT2_REAL - TOT2_EST,
         100.0 * (TOT2_REAL - TOT2_EST) / TOT2_EST, TOTA_EST, TOTA_REAL))

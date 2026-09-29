"""
Pruebas del módulo de Reportes — CU-42 a CU-46 (RF-41 a RF-45).

Lo que se verifica no es que las pantallas respondan 200, sino que los números
sean los correctos: un reporte que carga bien y suma mal es peor que uno que
falla, porque nadie lo revisa.
"""
import io
from datetime import date, timedelta
from decimal import Decimal

from django.test import Client, TestCase
from django.urls import reverse

from adquisiciones.models import OrdenCompra, OrdenCompraLinea
from inventario.models import Material, MovimientoInventario
from proveedores.models import Proveedor
from proyectos.models import Itemizado, Proyecto
from solicitudes.models import SolicitudDetalle, SolicitudMaterial
from usuarios.models import Usuario

from .consultas import (compras_por_proveedor, consumo_por_proyecto,
                        desviacion_presupuestaria, mermas_y_perdidas,
                        proyectos_visibles)
from .exportar import FormatoNoSoportado, exportar


class BaseReportes(TestCase):
    """Un proyecto con la cadena completa, para que los reportes tengan qué contar."""

    def setUp(self):
        self.client = Client()
        self.admin = Usuario.objects.create_user(
            "adm_rep", password="clave12345", email="adm@m5.cl", rol="ADMIN")
        self.ea = Usuario.objects.create_user(
            "ea_rep", password="clave12345", email="ea@m5.cl",
            rol="ENCARGADO_ADQUISICIONES")
        self.jefe = Usuario.objects.create_user(
            "jp_rep", password="clave12345", email="jp@m5.cl", rol="JEFE_PROYECTO")
        self.otro_jefe = Usuario.objects.create_user(
            "jp2_rep", password="clave12345", email="jp2@m5.cl", rol="JEFE_PROYECTO")
        self.bodeguero = Usuario.objects.create_user(
            "bod_rep", password="clave12345", email="bod@m5.cl", rol="BODEGUERO")

        self.proyecto = Proyecto.objects.create(
            nombre="Torre Alfa", mandante="Inmobiliaria X",
            fecha_inicio=date(2026, 6, 1), jefe_proyecto=self.jefe)
        self.ajeno = Proyecto.objects.create(
            nombre="Torre Beta", mandante="Inmobiliaria Y",
            fecha_inicio=date(2026, 6, 1), jefe_proyecto=self.otro_jefe)

        self.partida = Itemizado.objects.create(
            proyecto=self.proyecto, codigo_partida="OG-01", descripcion="Hormigón",
            unidad_medida="m3", cant_presupuestada=100, cant_ejecutada=120)  # sobregirada

        self.cemento = Material.objects.create(
            nombre="Cemento", unidad_medida="saco", precio_referencia=5000,
            stock_actual=100, tipo=Material.Tipo.CONSUMIBLE)
        self.fierro = Material.objects.create(
            nombre="Fierro 12mm", unidad_medida="barra", precio_referencia=9000,
            stock_actual=50, tipo=Material.Tipo.CONSUMIBLE)

        self.sm = SolicitudMaterial.objects.create(
            proyecto=self.proyecto, emisor=self.ea,
            estado=SolicitudMaterial.Estado.APROBADA)
        SolicitudDetalle.objects.create(
            solicitud=self.sm, material=self.cemento, cantidad_solicitada=100,
            unidad_medida="saco", valor_unitario=5000, partida=self.partida)
        SolicitudDetalle.objects.create(
            solicitud=self.sm, material=self.fierro, cantidad_solicitada=20,
            unidad_medida="barra", valor_unitario=9000, partida=self.partida)

        # 40 sacos salieron a la obra: quedan 60 pendientes
        MovimientoInventario.objects.create(
            tipo=MovimientoInventario.Tipo.SALIDA, material=self.cemento,
            cantidad=40, proyecto=self.proyecto, registrado_por=self.bodeguero)
        # una merma de 10 sacos = $50.000
        MovimientoInventario.objects.create(
            tipo=MovimientoInventario.Tipo.MERMA, material=self.cemento,
            cantidad=10, proyecto=self.proyecto,
            motivo=MovimientoInventario.MotivoMerma.DANO,
            registrado_por=self.bodeguero)

        self.proveedor = Proveedor.objects.create(
            nombre="Ferretería Andes", rut="11111111-1", correo="p@x.cl")
        self.orden = OrdenCompra.objects.create(
            solicitud=self.sm, proveedor=self.proveedor, creada_por=self.ea,
            estado=OrdenCompra.Estado.ENVIADA)
        OrdenCompraLinea.objects.create(
            orden=self.orden, material=self.cemento, descripcion="Cemento",
            cantidad=100, unidad_medida="saco", valor_unitario=5000)
        # Un borrador que NO debe contarse como compra
        self.borrador = OrdenCompra.objects.create(
            solicitud=self.sm, proveedor=self.proveedor, creada_por=self.ea,
            estado=OrdenCompra.Estado.BORRADOR)
        OrdenCompraLinea.objects.create(
            orden=self.borrador, material=self.fierro, descripcion="Fierro",
            cantidad=999, unidad_medida="barra", valor_unitario=9000)


# ==========================================================================
#  CU-42 (RF-41) — Consumo consolidado por proyecto
# ==========================================================================

class ConsumoCU42Test(BaseReportes):

    def test_separa_lo_solicitado_de_lo_despachado(self):
        r = consumo_por_proyecto(self.proyecto)
        fila = next(f for f in r["filas"] if f["material"] == "Cemento")
        self.assertEqual(fila["solicitado"], Decimal("100"))
        self.assertEqual(fila["despachado"], Decimal("40"))
        self.assertEqual(fila["pendiente"], Decimal("60"))

    def test_calcula_costo_unitario_y_total(self):
        r = consumo_por_proyecto(self.proyecto)
        fila = next(f for f in r["filas"] if f["material"] == "Cemento")
        self.assertEqual(fila["costo_unitario"], Decimal("5000"))
        self.assertEqual(fila["costo_total"], Decimal("500000"))
        # 100×5000 + 20×9000
        self.assertEqual(r["totales"]["costo_total"], Decimal("680000"))

    def test_trae_el_codigo_interno_del_catalogo(self):
        r = consumo_por_proyecto(self.proyecto)
        self.assertTrue(all(f["codigo"].startswith("MAT-") for f in r["filas"]))

    def test_un_proyecto_sin_solicitudes_lo_dice(self):
        r = consumo_por_proyecto(self.ajeno)
        self.assertTrue(r["vacio"])
        self.assertIn("todavía no tiene materiales", r["mensaje_vacio"])

    # --- Excepción 1: el Jefe de Proyecto sólo ve lo suyo ---

    def test_el_jefe_solo_ve_sus_proyectos(self):
        visibles = proyectos_visibles(self.jefe)
        self.assertIn(self.proyecto, visibles)
        self.assertNotIn(self.ajeno, visibles)

    def test_adquisiciones_ve_todos_los_proyectos(self):
        visibles = proyectos_visibles(self.ea)
        self.assertIn(self.proyecto, visibles)
        self.assertIn(self.ajeno, visibles)

    def test_el_jefe_no_alcanza_un_proyecto_ajeno_por_la_url(self):
        self.client.login(username="jp_rep", password="clave12345")
        resp = self.client.get(reverse("reportes:consumo"), {"proyecto": self.ajeno.pk})
        self.assertEqual(resp.status_code, 200)
        self.assertIsNone(resp.context["reporte"])
        self.assertContains(resp, "no está entre los que tienes asignados")

    def test_la_pantalla_muestra_el_reporte(self):
        self.client.login(username="ea_rep", password="clave12345")
        resp = self.client.get(reverse("reportes:consumo"), {"proyecto": self.proyecto.pk})
        self.assertContains(resp, "Cemento")
        self.assertContains(resp, "Torre Alfa")


# ==========================================================================
#  CU-43 (RF-42) — Desviación presupuestaria
# ==========================================================================

class DesviacionCU43Test(BaseReportes):

    def test_compara_presupuestado_contra_consumido(self):
        r = desviacion_presupuestaria(self.proyecto)
        fila = r["filas"][0]
        self.assertEqual(fila["presupuestado"], Decimal("100"))
        self.assertEqual(fila["consumido"], Decimal("120"))
        self.assertEqual(fila["desviacion"], Decimal("20"))
        self.assertTrue(fila["sobregirada"])

    def test_el_porcentaje_de_desviacion_es_correcto(self):
        r = desviacion_presupuestaria(self.proyecto)
        self.assertEqual(round(float(r["filas"][0]["desviacion_pct"]), 1), 20.0)

    def test_lo_comprado_excluye_borradores(self):
        """Un borrador no es una compra: 100 del cemento, no 1099."""
        r = desviacion_presupuestaria(self.proyecto)
        self.assertEqual(r["filas"][0]["comprado"], Decimal("100"))

    def test_incluye_el_grafico_que_pide_el_requisito(self):
        r = desviacion_presupuestaria(self.proyecto)
        self.assertIn("grafico", r)
        self.assertEqual(len(r["grafico"]["barras"]), 1)
        barra = r["grafico"]["barras"][0]
        self.assertTrue(0 <= barra["pct_presupuestado"] <= 100)
        self.assertTrue(barra["sobregirada"])

    # --- Excepción 1: sin itemizado no hay contra qué comparar ---

    def test_sin_itemizado_avisa_en_vez_de_mostrar_tabla_vacia(self):
        r = desviacion_presupuestaria(self.ajeno)
        self.assertTrue(r["vacio"])
        self.assertTrue(r["sin_datos_de_entrada"])
        self.assertIn("no tiene itemizado cargado", r["mensaje_vacio"])

    def test_la_pantalla_ofrece_ir_a_cargar_el_itemizado(self):
        self.client.login(username="ea_rep", password="clave12345")
        resp = self.client.get(reverse("reportes:desviacion"), {"proyecto": self.ajeno.pk})
        self.assertContains(resp, "cargar el itemizado")


# ==========================================================================
#  CU-44 (RF-43) — Mermas y pérdidas
# ==========================================================================

class MermasCU44Test(BaseReportes):

    def test_valoriza_la_perdida_en_pesos(self):
        r = mermas_y_perdidas()
        self.assertEqual(len(r["filas"]), 1)
        self.assertEqual(r["filas"][0]["valorizacion"], Decimal("50000"))  # 10 × 5000
        self.assertEqual(r["totales"]["valorizacion"], Decimal("50000"))

    def test_filtra_por_proyecto(self):
        self.assertEqual(len(mermas_y_perdidas(proyecto=self.proyecto)["filas"]), 1)
        self.assertEqual(len(mermas_y_perdidas(proyecto=self.ajeno)["filas"]), 0)

    def test_filtra_por_rango_de_fechas(self):
        hoy = date.today()
        self.assertEqual(len(mermas_y_perdidas(desde=hoy, hasta=hoy)["filas"]), 1)
        manana = hoy + timedelta(days=1)
        self.assertEqual(len(mermas_y_perdidas(desde=manana)["filas"]), 0)

    def test_solo_cuenta_mermas_no_salidas(self):
        """La salida a obra de 40 sacos no es una pérdida."""
        r = mermas_y_perdidas()
        self.assertEqual(len(r["filas"]), 1)
        self.assertEqual(r["filas"][0]["cantidad"], Decimal("10"))

    def test_traduce_el_motivo_a_texto_legible(self):
        r = mermas_y_perdidas()
        self.assertEqual(r["filas"][0]["motivo"], "Daño")

    # --- Excepción 1 ---

    def test_sin_mermas_muestra_mensaje_no_tabla_vacia(self):
        MovimientoInventario.objects.filter(
            tipo=MovimientoInventario.Tipo.MERMA).delete()
        r = mermas_y_perdidas()
        self.assertTrue(r["vacio"])
        self.assertIn("buena noticia", r["mensaje_vacio"])


# ==========================================================================
#  CU-45 (RF-44) — Historial de compras por proveedor
# ==========================================================================

class ComprasCU45Test(BaseReportes):

    def test_agrupa_por_proveedor_con_volumen_y_monto(self):
        r = compras_por_proveedor()
        self.assertEqual(len(r["filas"]), 1)
        fila = r["filas"][0]
        self.assertEqual(fila["proveedor"], "Ferretería Andes")
        self.assertEqual(fila["transacciones"], 1)      # el borrador no cuenta
        self.assertEqual(fila["monto_total"], Decimal("500000"))

    def test_un_borrador_no_es_una_compra(self):
        r = compras_por_proveedor()
        self.assertNotEqual(r["filas"][0]["monto_total"], Decimal("9491000"))

    def test_calcula_el_promedio_por_orden(self):
        r = compras_por_proveedor()
        self.assertEqual(r["filas"][0]["ticket_promedio"], Decimal("500000"))

    def test_filtra_por_periodo(self):
        manana = date.today() + timedelta(days=1)
        self.assertTrue(compras_por_proveedor(desde=manana)["vacio"])

    def test_solo_adquisiciones_y_administracion_entran(self):
        self.client.login(username="jp_rep", password="clave12345")
        self.assertEqual(self.client.get(reverse("reportes:compras")).status_code, 403)
        self.client.login(username="ea_rep", password="clave12345")
        self.assertEqual(self.client.get(reverse("reportes:compras")).status_code, 200)

    # --- Excepción 1 ---

    def test_sin_compras_en_el_periodo_lo_dice(self):
        manana = (date.today() + timedelta(days=1)).isoformat()
        self.client.login(username="ea_rep", password="clave12345")
        resp = self.client.get(reverse("reportes:compras"), {"desde": manana})
        self.assertContains(resp, "No hay órdenes de compra emitidas")


# ==========================================================================
#  CU-46 (RF-45) — Exportando reportes a Excel o PDF
# ==========================================================================

class ExportacionCU46Test(BaseReportes):

    def setUp(self):
        super().setUp()
        self.client.login(username="ea_rep", password="clave12345")

    def _descargar(self, nombre_url, formato, **params):
        params["formato"] = formato
        return self.client.get(reverse(f"reportes:{nombre_url}"), params)

    def test_los_cuatro_reportes_se_exportan_a_xlsx(self):
        from openpyxl import load_workbook
        casos = [("consumo", {"proyecto": self.proyecto.pk}),
                 ("desviacion", {"proyecto": self.proyecto.pk}),
                 ("mermas", {}), ("compras", {})]
        for nombre, params in casos:
            with self.subTest(reporte=nombre):
                resp = self._descargar(nombre, "xlsx", **params)
                self.assertEqual(resp.status_code, 200)
                self.assertIn("spreadsheetml", resp["Content-Type"])
                self.assertIn(".xlsx", resp["Content-Disposition"])
                libro = load_workbook(io.BytesIO(resp.content))
                self.assertGreater(libro.active.max_row, 5)

    def test_los_cuatro_reportes_se_exportan_a_pdf(self):
        casos = [("consumo", {"proyecto": self.proyecto.pk}),
                 ("desviacion", {"proyecto": self.proyecto.pk}),
                 ("mermas", {}), ("compras", {})]
        for nombre, params in casos:
            with self.subTest(reporte=nombre):
                resp = self._descargar(nombre, "pdf", **params)
                self.assertEqual(resp.status_code, 200)
                self.assertEqual(resp["Content-Type"], "application/pdf")
                self.assertTrue(resp.content.startswith(b"%PDF"))

    def test_el_excel_trae_los_datos_del_reporte(self):
        from openpyxl import load_workbook
        resp = self._descargar("consumo", "xlsx", proyecto=self.proyecto.pk)
        hoja = load_workbook(io.BytesIO(resp.content)).active
        texto = " ".join(str(c.value) for fila in hoja.iter_rows() for c in fila if c.value)
        self.assertIn("Cemento", texto)
        self.assertIn("Torre Alfa", texto)

    def test_el_nombre_del_archivo_identifica_el_reporte(self):
        resp = self._descargar("mermas", "xlsx")
        self.assertIn("mermas", resp["Content-Disposition"].lower())

    # --- Excepción 1: hay que elegir un formato válido ---

    def test_un_formato_invalido_no_genera_archivo(self):
        reporte = mermas_y_perdidas()
        for malo in ("word", "csv", "", None, "docx"):
            with self.subTest(formato=malo):
                with self.assertRaises(FormatoNoSoportado):
                    exportar(reporte, malo)

    def test_la_vista_explica_que_el_formato_no_sirve(self):
        resp = self.client.get(reverse("reportes:mermas"), {"formato": "word"}, follow=True)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "formato válido")

    def test_un_reporte_vacio_tambien_se_exporta(self):
        """Poder llevarse la evidencia de que no hubo mermas también sirve."""
        MovimientoInventario.objects.filter(
            tipo=MovimientoInventario.Tipo.MERMA).delete()
        resp = self._descargar("mermas", "pdf")
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.content.startswith(b"%PDF"))


class PanelReportesTest(BaseReportes):

    def test_el_panel_lista_los_reportes(self):
        self.client.login(username="ea_rep", password="clave12345")
        resp = self.client.get(reverse("reportes:panel"))
        self.assertContains(resp, "Consumo consolidado")
        self.assertContains(resp, "Desviación presupuestaria")
        self.assertContains(resp, "Mermas y pérdidas")
        self.assertContains(resp, "Historial de compras")

    def test_al_jefe_no_se_le_ofrece_el_historial_de_compras(self):
        self.client.login(username="jp_rep", password="clave12345")
        resp = self.client.get(reverse("reportes:panel"))
        self.assertNotContains(resp, reverse("reportes:compras"))

    def test_bodega_no_entra_a_reportes(self):
        self.client.login(username="bod_rep", password="clave12345")
        self.assertEqual(self.client.get(reverse("reportes:panel")).status_code, 403)

    def test_contabilidad_tampoco(self):
        """RF-41 a RF-44 nombran Jefe de Proyecto, Adquisiciones y Administración."""
        contador = Usuario.objects.create_user(
            "cont_rep", password="clave12345", email="c@m5.cl", rol="CONTABILIDAD")
        self.client.login(username="cont_rep", password="clave12345")
        self.assertEqual(self.client.get(reverse("reportes:panel")).status_code, 403)

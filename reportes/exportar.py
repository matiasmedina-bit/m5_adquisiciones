"""
CU-46 (RF-45) — Exportando reportes a Excel o PDF.

Una sola implementación para los cuatro reportes. Es posible porque todos
salen de `consultas.py` con la misma forma (columnas + filas + totales): el
exportador no sabe si está imprimiendo mermas o compras por proveedor, sólo
lee la estructura. Agregar un quinto reporte no toca este archivo.

Formatos soportados: XLSX (openpyxl) y PDF (reportlab), los dos que nombra el
requisito. Cualquier otro valor se rechaza — Excepción 1 del caso de uso.
"""
import io
from decimal import Decimal
from pathlib import Path

from django.conf import settings
from django.utils import timezone

FORMATOS = ("xlsx", "pdf")

AZUL = "1F3864"
AZUL_CLARO = "2E5FA3"
GRIS = "F2F2F2"


class FormatoNoSoportado(ValueError):
    """Excepción 1 del CU-46: hay que elegir entre XLSX y PDF."""


class LibreriaFaltante(RuntimeError):
    pass


# --------------------------------------------------------------------------
# Formato de valores. Chile: miles con punto, decimales con coma.
# --------------------------------------------------------------------------

def _miles(valor):
    try:
        return f"{int(valor):,}".replace(",", ".")
    except (TypeError, ValueError):
        return "0"


def _numero(valor):
    try:
        numero = float(valor)
    except (TypeError, ValueError):
        return str(valor)
    if numero == int(numero):
        return _miles(int(numero))
    return f"{numero:,.2f}".replace(",", "·").replace(".", ",").replace("·", ".")


def formatear(valor, tipo):
    if valor is None:
        return ""
    if tipo == "moneda":
        return "$" + _miles(valor)
    if tipo == "numero":
        return _numero(valor)
    return str(valor)


def _crudo(valor):
    """Valor sin formato, para que Excel lo trate como número y no como texto."""
    if isinstance(valor, Decimal):
        return float(valor)
    return valor


def nombre_archivo(reporte, formato):
    from django.utils.text import slugify
    base = slugify(reporte["titulo"]) or "reporte"
    sub = slugify(reporte.get("subtitulo", ""))
    sello = timezone.localtime().strftime("%Y%m%d")
    partes = [p for p in (base, sub, sello) if p]
    return "-".join(partes) + "." + formato


# ==========================================================================
#  XLSX
# ==========================================================================

def a_xlsx(reporte) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise LibreriaFaltante(
            "Falta openpyxl para exportar a Excel. Instálalo con: pip install openpyxl"
        ) from exc

    libro = Workbook()
    hoja = libro.active
    hoja.title = reporte["titulo"][:31]

    blanco = Font(color="FFFFFF", bold=True, size=11)
    relleno = PatternFill("solid", fgColor=AZUL)
    borde = Border(bottom=Side(style="thin", color="BFBFBF"))

    fila = 1
    # --- Cabecera ---
    hoja.cell(row=fila, column=1, value=reporte["titulo"]).font = Font(bold=True, size=14, color=AZUL)
    fila += 1
    if reporte.get("subtitulo"):
        hoja.cell(row=fila, column=1, value=reporte["subtitulo"]).font = Font(size=11, color="595959")
        fila += 1
    hoja.cell(row=fila, column=1,
              value=f"Constructora M5 SpA · generado el "
                    f"{timezone.localtime().strftime('%d/%m/%Y %H:%M')}").font = Font(size=9, italic=True, color="808080")
    fila += 2

    # --- Filtros aplicados: sin esto una planilla suelta no dice de qué es ---
    for etiqueta, valor in reporte.get("filtros", []):
        hoja.cell(row=fila, column=1, value=f"{etiqueta}:").font = Font(bold=True, size=9)
        hoja.cell(row=fila, column=2, value=str(valor)).font = Font(size=9)
        fila += 1
    fila += 1

    columnas = reporte["columnas"]
    if not columnas:
        hoja.cell(row=fila, column=1, value=reporte.get("mensaje_vacio", "Sin datos."))
        buffer = io.BytesIO(); libro.save(buffer); return buffer.getvalue()

    # --- Encabezados ---
    fila_encabezado = fila
    for col, columna in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila, column=col, value=columna["titulo"])
        celda.font = blanco
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center", vertical="center")
    fila += 1

    # --- Datos ---
    for registro in reporte["filas"]:
        for col, columna in enumerate(columnas, start=1):
            valor = registro.get(columna["clave"])
            celda = hoja.cell(row=fila, column=col, value=_crudo(valor))
            celda.border = borde
            if columna["tipo"] == "moneda":
                celda.number_format = '"$"#,##0'
                celda.alignment = Alignment(horizontal="right")
            elif columna["tipo"] == "numero":
                celda.number_format = "#,##0.##"
                celda.alignment = Alignment(horizontal="right")
        fila += 1

    # --- Totales ---
    totales = reporte.get("totales") or {}
    if totales:
        for col, columna in enumerate(columnas, start=1):
            if columna["clave"] in totales:
                celda = hoja.cell(row=fila, column=col, value=_crudo(totales[columna["clave"]]))
                celda.font = Font(bold=True)
                celda.fill = PatternFill("solid", fgColor=GRIS)
                if columna["tipo"] == "moneda":
                    celda.number_format = '"$"#,##0'
                elif columna["tipo"] == "numero":
                    celda.number_format = "#,##0.##"
            else:
                hoja.cell(row=fila, column=col).fill = PatternFill("solid", fgColor=GRIS)

    # --- Ancho de columna según el contenido real ---
    for col, columna in enumerate(columnas, start=1):
        largos = [len(columna["titulo"])]
        for registro in reporte["filas"][:200]:
            largos.append(len(formatear(registro.get(columna["clave"]), columna["tipo"])))
        hoja.column_dimensions[get_column_letter(col)].width = min(max(largos) + 4, 45)

    hoja.freeze_panes = hoja.cell(row=fila_encabezado + 1, column=1)
    hoja.auto_filter.ref = (f"A{fila_encabezado}:"
                            f"{get_column_letter(len(columnas))}{fila_encabezado + len(reporte['filas'])}")

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


# ==========================================================================
#  PDF
# ==========================================================================

def a_pdf(reporte) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (Image, Paragraph, SimpleDocTemplate,
                                        Spacer, Table, TableStyle)
    except ImportError as exc:
        raise LibreriaFaltante(
            "Falta reportlab para exportar a PDF. Instálalo con: pip install reportlab"
        ) from exc

    columnas = reporte["columnas"]
    # Con muchas columnas el vertical no alcanza: la tabla se corta o se
    # comprime hasta ser ilegible. Se decide la orientación por el reporte.
    tamano = landscape(A4) if len(columnas) > 6 else A4

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=tamano,
        leftMargin=14 * mm, rightMargin=14 * mm,
        topMargin=14 * mm, bottomMargin=16 * mm,
        title=reporte["titulo"], author="Constructora M5 SpA",
    )
    estilos = getSampleStyleSheet()
    h1 = ParagraphStyle("h1m5", parent=estilos["Heading1"], fontSize=15,
                        textColor=colors.HexColor("#" + AZUL), spaceAfter=2)
    sub = ParagraphStyle("subm5", parent=estilos["Normal"], fontSize=10.5,
                         textColor=colors.HexColor("#595959"), spaceAfter=1)
    pie = ParagraphStyle("piem5", parent=estilos["Normal"], fontSize=8,
                         textColor=colors.HexColor("#808080"))
    celda = ParagraphStyle("celdam5", parent=estilos["Normal"], fontSize=7.5, leading=9)
    celda_cab = ParagraphStyle("cabm5", parent=celda, textColor=colors.white,
                               fontName="Helvetica-Bold")

    elementos = []

    logo = Path(settings.BASE_DIR) / "static" / "img" / "Logo-M5.png"
    if logo.exists():
        try:
            elementos.append(Image(str(logo), width=30 * mm, height=13 * mm, kind="proportional"))
            elementos.append(Spacer(1, 4))
        except Exception:
            pass

    elementos.append(Paragraph(reporte["titulo"], h1))
    if reporte.get("subtitulo"):
        elementos.append(Paragraph(reporte["subtitulo"], sub))
    elementos.append(Paragraph(
        f"Constructora M5 SpA · generado el "
        f"{timezone.localtime().strftime('%d/%m/%Y a las %H:%M')}", pie))
    elementos.append(Spacer(1, 8))

    filtros = reporte.get("filtros") or []
    if filtros:
        datos = [[Paragraph(f"<b>{e}</b>", celda), Paragraph(str(v), celda)] for e, v in filtros]
        tabla_filtros = Table(datos, colWidths=[35 * mm, 60 * mm], hAlign="LEFT")
        tabla_filtros.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F7F8FA")),
            ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#D5D9DE")),
            ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E6E9ED")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        elementos.append(tabla_filtros)
        elementos.append(Spacer(1, 10))

    if reporte.get("vacio") or not columnas:
        elementos.append(Paragraph(
            reporte.get("mensaje_vacio", "No hay datos para los filtros seleccionados."),
            estilos["Normal"]))
        doc.build(elementos)
        return buffer.getvalue()

    encabezado = [Paragraph(c["titulo"], celda_cab) for c in columnas]
    cuerpo = [encabezado]
    for registro in reporte["filas"]:
        cuerpo.append([
            Paragraph(formatear(registro.get(c["clave"]), c["tipo"]), celda)
            for c in columnas
        ])

    totales = reporte.get("totales") or {}
    if totales:
        cuerpo.append([
            Paragraph(f"<b>{formatear(totales.get(c['clave']), c['tipo'])}</b>", celda)
            if c["clave"] in totales else Paragraph("", celda)
            for c in columnas
        ])

    ancho_util = tamano[0] - 28 * mm
    tabla = Table(cuerpo, colWidths=[ancho_util / len(columnas)] * len(columnas),
                  repeatRows=1, hAlign="LEFT")
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + AZUL)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D9DDE3")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F8FA")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    # Alinear a la derecha lo que es número, como en cualquier planilla
    for indice, columna in enumerate(columnas):
        if columna["tipo"] in ("numero", "moneda"):
            estilo.append(("ALIGN", (indice, 1), (indice, -1), "RIGHT"))
    if totales:
        estilo.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E9ECEF")))
    tabla.setStyle(TableStyle(estilo))
    elementos.append(tabla)

    elementos.append(Spacer(1, 8))
    elementos.append(Paragraph(
        f"{len(reporte['filas'])} registro{'s' if len(reporte['filas']) != 1 else ''} · "
        f"Sistema de Gestión de Adquisiciones e Inventario", pie))

    doc.build(elementos)
    return buffer.getvalue()


# ==========================================================================

def exportar(reporte, formato):
    """
    Devuelve (bytes, nombre_archivo, content_type).

    Excepción 1 del CU-46: si el formato no es uno de los dos soportados, no se
    genera nada y se levanta FormatoNoSoportado para que la vista lo explique.
    """
    formato = (formato or "").lower().strip()
    if formato not in FORMATOS:
        raise FormatoNoSoportado(
            "Elige un formato válido para exportar: Excel (XLSX) o PDF.")
    if formato == "xlsx":
        return (a_xlsx(reporte), nombre_archivo(reporte, "xlsx"),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    return a_pdf(reporte), nombre_archivo(reporte, "pdf"), "application/pdf"
